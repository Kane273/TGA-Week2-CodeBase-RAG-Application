from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, Protocol

from docx import Document as DocxDocument
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class DocumentParsingError(ValueError):
    """Raised when a parser cannot safely extract text from a document."""


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    parser: str


class DocumentParser(Protocol):
    """Provider-neutral boundary for local or managed document parsers."""

    def parse(self, path: Path) -> ParsedDocument: ...


class LocalTextLoader(BaseLoader):
    def __init__(self, path: Path):
        self.path = path

    def lazy_load(self) -> Iterator[Document]:
        yield Document(page_content=self.path.read_text(encoding="utf-8"),
                       metadata={"parser": "utf-8-text", "source": str(self.path)})


class LocalPdfLoader(BaseLoader):
    def __init__(self, path: Path):
        self.path = path

    def lazy_load(self) -> Iterator[Document]:
        text = "\n\n".join(
            f"Page {number}\n{page.extract_text() or ''}"
            for number, page in enumerate(PdfReader(str(self.path)).pages, start=1)
        )
        yield Document(page_content=text,
                       metadata={"parser": "pypdf", "source": str(self.path)})


class LocalDocxLoader(BaseLoader):
    def __init__(self, path: Path):
        self.path = path

    def lazy_load(self) -> Iterator[Document]:
        document = DocxDocument(str(self.path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        yield Document(page_content=text,
                       metadata={"parser": "python-docx", "source": str(self.path)})


def build_document_loader(path: Path) -> BaseLoader:
    if path.suffix.lower() == ".pdf":
        return LocalPdfLoader(path)
    if path.suffix.lower() == ".docx":
        return LocalDocxLoader(path)
    return LocalTextLoader(path)


class LocalDocumentParser:
    """Extract text locally without credentials or external document transfer."""

    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = build_document_loader(path).load()[0]
            return ParsedDocument(document.page_content, document.metadata["parser"])
        except (UnicodeDecodeError, OSError, ValueError, PdfReadError) as error:
            raise DocumentParsingError(f"Unable to parse {path.name}") from error


def build_document_parser(provider: str) -> DocumentParser:
    if provider.lower() == "local":
        return LocalDocumentParser()
    raise ValueError(
        "DOCUMENT_PARSER_PROVIDER must be 'local'; additional providers must "
        "implement the DocumentParser interface"
    )
