from pathlib import Path

import pytest
from docx import Document as DocxDocument
from langchain_core.document_loaders import BaseLoader

from app.ingestion import (LANGUAGE_BY_EXTENSION, chunk_file, iter_source_files,
                           sensitive_content_reason)
from app.parsing import ParsedDocument, build_document_loader, build_document_parser


def test_chunks_include_source_lines(tmp_path: Path):
    source = tmp_path / "service.py"
    source.write_text("def answer():\n    return 42\n")
    docs = chunk_file(source, tmp_path, "demo", "main")
    assert docs
    assert docs[0].metadata["path"] == "service.py"
    assert docs[0].metadata["start_line"] == 1
    assert docs[0].metadata["end_line"] == 2
    assert docs[0].metadata["symbol_type"] in {"function_definition", "module_context"}


def test_ast_chunk_captures_symbol_name(tmp_path: Path):
    source = tmp_path / "service.py"
    source.write_text("class BillingService:\n    def charge_customer(self):\n        return True\n")
    docs = chunk_file(source, tmp_path, "demo", "main")
    assert any(doc.metadata["symbol"] == "BillingService" for doc in docs)


def test_ignored_directories_are_not_indexed(tmp_path: Path):
    ignored = tmp_path / "node_modules"
    ignored.mkdir()
    (ignored / "dependency.js").write_text("const secret = true")
    (tmp_path / "app.js").write_text("const app = true")
    assert [p.name for p in iter_source_files(tmp_path, 10)] == ["app.js"]


def test_secret_bearing_source_is_excluded(tmp_path: Path):
    source = tmp_path / "config.py"
    source.write_text('api_key = "this-is-a-secret-value"\n')
    assert chunk_file(source, tmp_path, "demo", "main") == []


def test_symlink_is_not_indexed(tmp_path: Path):
    outside = tmp_path / "outside.py"
    outside.write_text("safe = True")
    link = tmp_path / "linked.py"
    link.symlink_to(outside)
    assert link not in list(iter_source_files(tmp_path, 10))
    assert chunk_file(link, tmp_path, "demo", "main") == []


@pytest.mark.parametrize("content", [
    "token = 'ghp_abcdefghijklmnopqrstuvwxyz123456'",
    "Authorization: Bearer abcdefghijklmnopqrstuvwxyz.abcdefghijk",
    "DATABASE_URL=postgresql://admin:supersecret@example.com/app",
    "AccountKey=ABCDEFGHIJKLMNOPQRSTUVWXYZ1234567890+/=",
])
def test_expanded_sensitive_content_is_detected(content: str):
    assert sensitive_content_reason(content)


def test_high_entropy_assigned_credential_is_detected():
    assert sensitive_content_reason(
        'credential = "p8Jx2Qv9Lm4Nz7Kr5Wt0Bc6Hd3Fy1SaE"'
    ) == "high-entropy assigned credential"


def test_docx_document_is_extracted_and_chunked(tmp_path: Path):
    path = tmp_path / "architecture.docx"
    document = DocxDocument()
    document.add_heading("System Architecture")
    document.add_paragraph("The API sends indexing work to a background worker.")
    document.save(path)
    chunks = chunk_file(path, tmp_path, "docs", "uploads", "uploads")
    assert chunks
    assert "background worker" in chunks[0].page_content
    assert chunks[0].metadata["path"] == "uploads/architecture.docx"
    assert chunks[0].metadata["parser"] == "python-docx"


def test_chunking_accepts_a_provider_neutral_document_parser(tmp_path: Path):
    class StubParser:
        def parse(self, path: Path) -> ParsedDocument:
            return ParsedDocument("Text extracted by an optional parser.", "stub")

    path = tmp_path / "source.pdf"
    path.write_bytes(b"not-a-real-pdf")
    chunks = chunk_file(path, tmp_path, "docs", "main", parser=StubParser())
    assert chunks[0].metadata["parser"] == "stub"
    assert "optional parser" in chunks[0].page_content


def test_unknown_document_parser_provider_is_rejected():
    try:
        build_document_parser("managed-service")
    except ValueError as error:
        assert "DOCUMENT_PARSER_PROVIDER" in str(error)
    else:
        raise AssertionError("unknown parser provider should fail fast")


def test_javascript_family_uses_safe_text_chunking():
    assert ".js" not in LANGUAGE_BY_EXTENSION
    assert ".jsx" not in LANGUAGE_BY_EXTENSION
    assert ".ts" not in LANGUAGE_BY_EXTENSION
    assert ".tsx" not in LANGUAGE_BY_EXTENSION


def test_ast_parsing_can_be_disabled_for_untrusted_remote_code(tmp_path: Path):
    source = tmp_path / "remote.py"
    source.write_text("def remote_function():\n    return True\n")
    docs = chunk_file(source, tmp_path, "remote", "main", ast_enabled=False)
    assert docs
    assert all(doc.metadata["symbol_type"] == "text_chunk" for doc in docs)


def test_local_files_use_langchain_document_loaders(tmp_path: Path):
    source = tmp_path / "README.md"
    source.write_text("Local LangChain loader")
    loader = build_document_loader(source)
    assert isinstance(loader, BaseLoader)
    assert loader.load()[0].page_content == "Local LangChain loader"
